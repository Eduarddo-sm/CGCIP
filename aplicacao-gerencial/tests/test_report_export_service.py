from __future__ import annotations

import io
import unittest
import zipfile

from openpyxl import load_workbook

from services.report_export_service import ReportExportService


class ReportExportServiceTestCase(unittest.TestCase):
    def setUp(self) -> None:
        self.service = ReportExportService()
        self.headers = ["CLIENTE", "VALOR"]
        self.rows = [{"CLIENTE": "Empresa & Filhos", "VALOR": 1250.5}]

    def test_csv_is_excel_compatible_and_uses_semicolon(self) -> None:
        content = self.service.csv_bytes(self.headers, self.rows, str)
        self.assertTrue(content.startswith(b"\xef\xbb\xbf"))
        self.assertIn(b"CLIENTE;VALOR", content)

    def test_xlsx_is_formatted_table_without_frozen_header(self) -> None:
        headers = ["CLIENTE", "DATA ACORDO", "DT AJUIZAMENTO", "VALOR DO ACORDO", "NPJ"]
        rows = [{
            "CLIENTE": "Empresa & Filhos",
            "DATA ACORDO": "2026-07-29",
            "DT AJUIZAMENTO": "2024-01-05",
            "VALOR DO ACORDO": "1250,50",
            "NPJ": "00123456789012",
        }]
        content = self.service.xlsx_bytes("Producao", headers, rows, str)
        with zipfile.ZipFile(io.BytesIO(content)) as workbook:
            self.assertIn("xl/worksheets/sheet1.xml", workbook.namelist())
            self.assertIn("xl/tables/table1.xml", workbook.namelist())
            sheet = workbook.read("xl/worksheets/sheet1.xml").decode("utf-8")
            self.assertNotIn("<pane", sheet)

        workbook = load_workbook(io.BytesIO(content), data_only=True)
        sheet = workbook["Producao"]
        self.assertEqual(len(sheet.tables), 1)
        self.assertEqual(sheet.freeze_panes, None)
        self.assertEqual(sheet["B2"].number_format, "dd/mm/yyyy")
        self.assertEqual(sheet["C2"].number_format, "dd/mm/yyyy")
        self.assertIn('"R$"* #,##0.00', sheet["D2"].number_format)
        self.assertEqual(sheet["D2"].value, 1250.5)
        self.assertEqual(sheet["E2"].value, "00123456789012")


if __name__ == "__main__":
    unittest.main()
